"""Template blueprint scanner for docx tables."""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Dict, List

from docx import Document


CHAPTER_IDS = {"4", "5", "6", "7", "8", "9", "10", "B"}


@dataclass
class TableBlueprint:
    table_id: str
    table_type: str
    table_index: int
    header_text: str = ""


class TemplateBlueprint:
    """Scans a docx template to build table metadata."""

    def __init__(self, tables: List[TableBlueprint]):
        self.tables = tables
        self.by_id: Dict[str, TableBlueprint] = {t.table_id: t for t in tables}
        self.by_type: Dict[str, List[TableBlueprint]] = {}
        for table in tables:
            self.by_type.setdefault(table.table_type, []).append(table)

    @classmethod
    def from_document(cls, doc: Document) -> "TemplateBlueprint":
        tables: List[TableBlueprint] = []
        for idx, table in enumerate(doc.tables):
            if not table.rows:
                continue

            first_cell = table.cell(0, 0).text.strip()
            if "\n" in first_cell:
                first_cell = first_cell.split("\n", 1)[0].strip()

            table_id = cls._normalize_table_id(first_cell)
            if not table_id:
                continue

            table_type = cls._classify_table(table_id, first_cell)
            tables.append(TableBlueprint(
                table_id=table_id,
                table_type=table_type,
                table_index=idx,
                header_text=first_cell
            ))

        return cls(tables)

    @staticmethod
    def _normalize_table_id(text: str) -> str:
        if not text:
            return ""

        text = " ".join(text.split())

        if "安全防護總攬表" in text:
            return "安全防護總攬表"
        if "能量源圖" in text:
            return "能量源圖"

        if text in CHAPTER_IDS:
            return text
        if len(text) == 1 and text.isalpha():
            return text

        parts = text.split(",")
        if parts:
            return parts[0].strip()

        return text

    @staticmethod
    def _classify_table(table_id: str, header_text: str) -> str:
        if table_id == "安全防護總攬表":
            return "overview"
        if table_id == "能量源圖":
            return "energy_diagram"
        if table_id in CHAPTER_IDS:
            return "chapter"

        if re.match(r"^[A-Z]\.?\d", table_id) or re.match(r"^\d+\.\d", table_id):
            return "appended"
        if table_id in {"X"}:
            return "appended"

        return "unknown"
